from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = r"D:\nextwave\output\Layer_2_Parent_Decision_Brief.docx"


def hyperlink(paragraph, label, url):
    part = paragraph.part
    rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rp = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "19548A")
    rp.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rp.append(u)
    r.append(rp)
    t = OxmlElement("w:t")
    t.text = label
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.76)
sec.bottom_margin = Inches(0.70)
sec.left_margin = Inches(0.84)
sec.right_margin = Inches(0.84)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.4)
normal.font.color.rgb = RGBColor(24, 30, 36)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.13

title = doc.styles["Title"]
title.font.name = "Aptos Display"
title.font.size = Pt(21)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(5)
title.paragraph_format.keep_with_next = True

head = doc.styles["Heading 2"]
head.font.name = "Aptos"
head.font.size = Pt(11.3)
head.font.bold = True
head.font.color.rgb = RGBColor(0, 0, 0)
head.paragraph_format.space_before = Pt(9)
head.paragraph_format.space_after = Pt(3)
head.paragraph_format.keep_with_next = True

small = doc.styles.add_style("Source Line", WD_STYLE_TYPE.PARAGRAPH)
small.base_style = normal
small.font.name = "Aptos"
small.font.size = Pt(8.3)
small.font.color.rgb = RGBColor(80, 87, 95)
small.paragraph_format.space_before = Pt(6)
small.paragraph_format.space_after = Pt(0)

p = doc.add_paragraph(style="Title")
p.add_run("NIAT Parent Decision Brief")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Track B  |  Lead to sale")
r.bold = True
r.font.size = Pt(9.5)
r.font.color.rgb = RGBColor(75, 83, 91)

doc.add_heading("The avenue", level=2)
p = doc.add_paragraph()
p.add_run(
    "After NAT and counselling, send each undecided family a private, shareable decision brief for the university discussed with them. "
    "It records who awards the degree, what NIAT separately provides, verified university and NIAT fees, scholarship terms, relevant hostel costs, and the family's open questions. "
    "Student and parent can each flag what remains unclear and book one joint call with the same counsellor. "
    "NIAT publicly offers a brochure, one-to-one counselling and WhatsApp contact, but I could not find this shared post-counselling record in its public funnel. "
    "My guess is that institution- and NAT-dependent costs make a standard follow-up document difficult to maintain."
)

doc.add_heading("Why it fits NIAT", level=2)
p = doc.add_paragraph()
p.add_run(
    "A family is choosing both a university degree and a separate NIAT upskilling program. The university controls admission and awards the degree; NIAT does neither. "
    "NIAT also says exact fees depend on the institution and NAT result and should be discussed after qualification. "
    "A counsellor-verified brief turns those variable facts into one record a parent can inspect before paying."
)

doc.add_heading("My hypothesis", level=2)
p = doc.add_paragraph()
p.add_run(
    "I expect the brief to reduce the delay between a student's interest and a parent's financial decision. "
    "My planning estimate is that 10% of qualified, counselled but unpaid leads make an NIAT upskilling enrollment payment within 14 days under usual follow-up; "
    "the brief and joint call could lift that to 16%. These are guesses, not NIAT performance data."
)

doc.add_heading("The one experiment I would run first", level=2)
p = doc.add_paragraph()
p.add_run(
    "In the next active intake, take 400 NAT-qualified leads who completed counselling, received a specific university suggestion, and have not paid. "
    "Randomly split them 200/200 within university and counselling week. One group gets usual follow-up; the other receives the brief within 24 hours and a joint-call link. "
    "A counsellor checks each fee, scholarship and university statement before sharing it. Measure both groups for 14 days from assignment."
)

doc.add_heading("How I would know it worked", level=2)
p = doc.add_paragraph()
p.add_run(
    "The metric is the share of assigned leads whose first NIAT upskilling enrollment payment appears in NIAT's payment records within 14 days. "
    "I would expand only if the brief group beats usual follow-up by at least 6 percentage points (for example, 16% versus 10%). "
    "With 200 per group, that is 12 additional payments; it is a pilot signal, not a precise estimate of long-term impact."
)

p = doc.add_paragraph(style="Source Line")
p.add_run("Public evidence: ")
hyperlink(p, "NIAT admissions and FAQ", "https://apply.niatindia.com/login-gnc/")
p.add_run("  |  ")
hyperlink(p, "NIAT program and degree facts", "https://www.niatindia.com/llm-info")

doc.core_properties.title = "NIAT Parent Decision Brief"
doc.core_properties.subject = "Layer 2 growth avenue"
doc.core_properties.author = ""
doc.save(OUT)
print(OUT)
